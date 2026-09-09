import os
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.database import get_db
from app.models import Book, BookPage
from app.auth import get_current_user

router = APIRouter()

STORAGE_DIR = "storage/scraped"
os.makedirs(STORAGE_DIR, exist_ok=True)

@router.get("/download/{b_id}")
async def download_scraped_book(b_id: int, db: Session = Depends(get_db), current_user = Depends(get_current_user)):
    book = db.query(Book).filter(Book.id == b_id).first()
    if not book:
        raise HTTPException(status_code=404, detail="Book not found")
        
    if book.scrape_status != "SCRAPED":
        raise HTTPException(status_code=400, detail="Book is not fully scraped yet")
        
    filename = f"book_{book.book_id}.xlsx"
    file_path = os.path.join(STORAGE_DIR, filename)
    
    # Generate on the fly if it doesn't exist, or regenerate it
    wb = Workbook()
    ws = wb.active
    ws.title = f"Book_{book.book_id}"[:31]
    
    headers = [
        "Book_ID", "Detail_URL", "Book_Title", "Language_Hint", "Author", "Publisher",
        "Total_Pages", "Page_Number", "Page_URL", "Page_Title", "Page_Text", "Page_Meta_Description",
        "Page_Has_Audio", "Page_Has_PDF", "Page_Has_Translation"
    ]
    
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        
    pages = db.query(BookPage).filter(BookPage.book_id == book.id).order_by(BookPage.page_number).all()
    
    # Detail Row
    detail_row = [
        book.book_id, book.detail_url, book.title, book.language, book.author, book.publisher,
        book.total_pages, "DETAIL", book.detail_url, book.title, "", "",
        False, False, False
    ]
    for c, v in enumerate(detail_row, 1):
        ws.cell(2, c, v)
        
    row = 3
    for p in pages:
        vals = [
            book.book_id, book.detail_url, book.title, book.language, book.author, book.publisher,
            book.total_pages, p.page_number, p.page_url, p.page_title, p.page_text, p.page_meta_description,
            p.page_has_audio, p.page_has_pdf, p.page_has_translation
        ]
        for c, v in enumerate(vals, 1):
            ws.cell(row, c, v)
        row += 1
        
    ws.freeze_panes = "A2"
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = min(55, max(14, len(headers[c - 1]) + 2))
        
    wb.save(file_path)
    
    return FileResponse(
        path=file_path, 
        filename=filename,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
